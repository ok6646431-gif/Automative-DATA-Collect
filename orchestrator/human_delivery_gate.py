"""User-facing delivery gate: collection PASS is NOT completed delivery.

JSON manifests/receipts are internal diagnostics, never the user deliverable.
Only an actual human-readable PDF/XLSX archive, accepted on the final ZIP bytes
and confirmed as an uploaded Actions artifact, may be called available.
"""
from __future__ import annotations

import argparse
import hashlib
import io
import json
import zipfile
from pathlib import Path

from orchestrator.archive_acceptance import validate_archive_zip

USER_PREFIX = "01_사용자자료/"
ALLOWED_USER_SUFFIXES = frozenset({".pdf", ".xlsx"})


def read_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def _checksum(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1048576), b""):
            h.update(chunk)
    return h.hexdigest()


def check_archive(zip_path: str | Path, summary_path: str | Path,
                  manifest_path: str | Path) -> dict:
    path = Path(zip_path)
    summary_file = Path(summary_path)
    manifest_file = Path(manifest_path)
    problems: list[str] = []
    if not path.is_file() or path.stat().st_size == 0:
        return {"status": "DELIVERY_NOT_READY", "problems": ["HUMAN_ZIP_NOT_CREATED"]}
    if not summary_file.is_file() or not manifest_file.is_file():
        return {"status": "DELIVERY_NOT_READY", "problems": ["ARCHIVE_AUDIT_MISSING"]}
    try:
        summary = read_json(summary_file)
        manifest = read_json(manifest_file)
    except (OSError, json.JSONDecodeError) as e:
        return {"status": "DELIVERY_NOT_READY", "problems": ["ARCHIVE_AUDIT_NOT_READABLE:" + str(e)]}
    env_qa = summary.get("envinfo_content_qa") or {}
    expected = env_qa.get("expected_items")
    if expected is not None:
        try:
            expected = int(expected)
        except (TypeError, ValueError):
            problems.append("ENVINFO_EXPECTED_COUNT_INVALID")
            expected = None
    if env_qa.get("pass") is False:
        problems.append("ENVINFO_CONTENT_QA_FAILED")
    if summary.get("sustainability_korean_delivery_qa", {}).get("pass") is False:
        problems.append("SUSTAINABILITY_KOREAN_QA_FAILED")
    checks = summary.get("acceptance_checks") or {}
    for key in ("envinfo_content_fidelity", "sustainability_korean_pdf_fidelity"):
        if checks.get(key) is False:
            problems.append("ARCHIVE_FIDELITY_FAILED:" + key)
    archived = manifest.get("human_archive") or {}
    if archived.get("status") != "PASS":
        problems.append("HUMAN_ARCHIVE_MANIFEST_NOT_PASS")
    if archived.get("zip_file") != path.name:
        problems.append("HUMAN_ARCHIVE_MANIFEST_ZIP_MISMATCH")
    if summary.get("zip_sha256") and summary["zip_sha256"] != _checksum(path):
        problems.append("HUMAN_ZIP_SHA256_MISMATCH")
    accepted = validate_archive_zip(path, expected_env_pdf_count=expected)
    if accepted["status"] != "PASS":
        problems.extend("USER_ZIP_INVALID:" + str(item.get("check"))
                        for item in accepted.get("failures", []))
    counts = {"PDF": 0, "XLSX": 0, "OTHER": 0}
    try:
        with zipfile.ZipFile(path) as zf:
            root = accepted.get("archive_root")
            if not root:
                roots = {n.split("/", 1)[0] for n in zf.namelist() if "/" in n}
                root = next(iter(roots)) if len(roots) == 1 else ""
            for item in zf.infolist():
                if item.is_dir() or not root or not item.filename.startswith(root + "/"):
                    continue
                rel = item.filename[len(root) + 1:]
                if not rel.startswith(USER_PREFIX):
                    continue
                suffix = Path(rel).suffix.lower()
                if suffix == ".pdf":
                    counts["PDF"] += 1
                elif suffix == ".xlsx":
                    counts["XLSX"] += 1
                else:
                    counts["OTHER"] += 1
    except (OSError, zipfile.BadZipFile, RuntimeError) as e:
        problems.append("USER_ZIP_UNREADABLE:" + str(e))
    # Structural signatures alone cannot prove that user documents can open.
    # Reparse the actual ZIP member bytes with PDF and XLSX readers. Archive
    # jobs already install both dependencies, so missing readers fail closed.
    try:
        from pypdf import PdfReader
        from openpyxl import load_workbook
        with zipfile.ZipFile(path) as zf:
            roots = {item.filename.split("/", 1)[0] for item in zf.infolist()
                     if not item.is_dir() and "/" in item.filename}
            archive_root = next(iter(roots)) if len(roots) == 1 else ""
            for item in zf.infolist():
                if item.is_dir() or not archive_root:
                    continue
                prefix = archive_root + "/" + USER_PREFIX
                if not item.filename.startswith(prefix):
                    continue
                suffix = Path(item.filename).suffix.lower()
                if suffix not in ALLOWED_USER_SUFFIXES:
                    continue
                try:
                    data = io.BytesIO(zf.read(item))
                    if suffix == ".pdf":
                        pdf = PdfReader(data, strict=False)
                        if pdf.is_encrypted or len(pdf.pages) < 1:
                            raise ValueError("PDF has no accessible pages")
                    else:
                        wb = load_workbook(data, read_only=True, data_only=False)
                        try:
                            if not wb.sheetnames:
                                raise ValueError("XLSX has no worksheets")
                        finally:
                            wb.close()
                except Exception as exc:
                    problems.append("USER_FILE_CANNOT_OPEN:" + item.filename +
                                    ":" + type(exc).__name__)
    except (ImportError, OSError, zipfile.BadZipFile) as exc:
        problems.append("HUMAN_FORMAT_READER_UNAVAILABLE:" + type(exc).__name__)
    if counts["XLSX"] < 4:
        problems.append("FOUR_PUBLIC_USER_XLSX_EXPORTS_REQUIRED")
    if not (counts["PDF"] or counts["XLSX"]):
        problems.append("NO_PDF_OR_EXCEL_USER_FILES")
    if counts["OTHER"]:
        problems.append("NON_PDF_XLSX_IN_USER_LAYER")
    completeness_file = manifest_file.parent / "Collection_Completeness.json"
    completeness = read_json(completeness_file) if completeness_file.exists() else {}
    limited = (
        manifest.get("validation") != "PASS"
        or completeness.get("status") != "COMPLETE"
    )
    return {
        "status": ("DELIVERY_NOT_READY" if problems else
                   "HUMAN_ARCHIVE_READY_WITH_LIMITATIONS" if limited else
                   "HUMAN_ARCHIVE_READY"),
        "problems": problems,
        "zip_path": str(path),
        "zip_bytes": path.stat().st_size,
        "zip_sha256": _checksum(path),
        "human_files": counts,
        "envinfo_body_expected": expected,
        "collection_validation": manifest.get("validation", "MISSING"),
        "collection_completeness": completeness.get("status", "UNKNOWN"),
        "human_archive_acceptance": accepted.get("status"),
        "notice": ("JSON and collector success are internal evidence only; "
                   "only a validated downloadable ZIP with user-facing PDF/XLSX "
                   "can be announced as ready."),
    }


def confirm_publication(report: dict, artifact_id: str, run_url: str) -> dict:
    """GitHub upload confirmation is separate from ZIP validation."""
    ready = report.get("status") in {
        "HUMAN_ARCHIVE_READY", "HUMAN_ARCHIVE_READY_WITH_LIMITATIONS"
    }
    valid_id = str(artifact_id or "").strip().isdigit() and int(artifact_id) > 0
    if not ready or not valid_id or not str(run_url).startswith("https://github.com/"):
        return {
            "status": "DELIVERY_NOT_READY",
            "reason": "VALIDATED_HUMAN_ZIP_OR_CONFIRMED_UPLOAD_MISSING",
            "human_zip_ready": ready,
            "artifact_id_present": valid_id,
        }
    return {
        "status": ("USER_ARCHIVE_AVAILABLE" if report["status"] == "HUMAN_ARCHIVE_READY"
                   else "USER_ARCHIVE_AVAILABLE_WITH_LIMITATIONS"),
        "artifact_id": str(artifact_id),
        "run_url": run_url,
        "file_type": "ZIP_CONTAINING_PDF_AND_XLSX",
        "human_files": report.get("human_files"),
        "limitations": (report.get("collection_validation") != "PASS"
                        or report.get("collection_completeness") != "COMPLETE"),
        "message": (
            "PDF/XLSX ZIP published; JSON diagnostics are not a deliverable. "
            "Collection-level PASS alone never implies a user-facing delivery."
        ),
    }


def main() -> int:
    p = argparse.ArgumentParser()
    p.add_argument("--zip", default="assembled/Human_Archive.zip")
    p.add_argument("--summary", default="assembled/Archive_Summary.json")
    p.add_argument("--manifest", default="assembled/Master_Manifest.json")
    p.add_argument("--out", default="assembled/Human_Delivery_Gate.json")
    p.add_argument("--artifact-id", default="")
    p.add_argument("--run-url", default="")
    p.add_argument("--confirm-upload", action="store_true")
    args = p.parse_args()
    target = Path(args.out)
    target.parent.mkdir(parents=True, exist_ok=True)
    if args.confirm_upload:
        ready = read_json(target)
        report = confirm_publication(ready, args.artifact_id, args.run_url)
    else:
        report = check_archive(args.zip, args.summary, args.manifest)
    target.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n",
                      encoding="utf-8")
    print(json.dumps({"status": report.get("status"),
                      "human_files": report.get("human_files", {}),
                      "problems": report.get("problems", [])}, ensure_ascii=False))
    return 0 if report["status"] in {
        "HUMAN_ARCHIVE_READY", "HUMAN_ARCHIVE_READY_WITH_LIMITATIONS",
        "USER_ARCHIVE_AVAILABLE", "USER_ARCHIVE_AVAILABLE_WITH_LIMITATIONS",
    } else 2


if __name__ == "__main__":
    raise SystemExit(main())
