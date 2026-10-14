"""Company-agnostic five-registry human delivery (PDF + Excel only).

Input: one `output/<SOURCE>/` per original public collector, optionally a ZIP
containing source-native official ESG PDF files. These are *not* JSON deliverables.
Internal JSON QA stays beside the ZIP, never inside it.

Unlike a report-only/special-case script, this module is usable by S-OIL,
ISU PETASYS, and future legal entities without renaming source-specific files.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import re
import shutil
import tempfile
import zipfile
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader

from orchestrator.archive_builder import (
    dict_rows_to_xlsx, read_csv, read_jsonl, render_envinfo_static_pdf,
)
from orchestrator.environment_evidence_review import review as review_environment_evidence

SOURCES = ("ENVINFO", "PRTR", "CHEM_STATS", "CLEANSYS_AIR", "SOOSIRO_WATER")
SOURCE_URLS = {
    "ENVINFO": "https://www.env-info.kr/",
    "PRTR": "https://icis.me.go.kr/prtr/",
    "CHEM_STATS": "https://icis.me.go.kr/",
    "CLEANSYS_AIR": "https://www.cleansys.or.kr/",
    "SOOSIRO_WATER": "https://www.soosiro.or.kr/",
}
USER_EXTENSIONS = {".pdf", ".xlsx"}
ACCEPTED_SOURCE_STATUSES = {
    "ENVINFO": {"DATA_FOUND", "NO_DATA_CONFIRMED", "NO_PUBLIC_RECORD_FOUND_IN_SCOPE"},
    "PRTR": {"DATA_FOUND", "NO_DATA_CONFIRMED"},
    "CHEM_STATS": {"DATA_FOUND", "NO_DATA_CONFIRMED"},
    "CLEANSYS_AIR": {"DATA_FOUND", "NO_DATA_CONFIRMED", "NO_FACILITY_MATCH_CONFIRMED"},
    "SOOSIRO_WATER": {"DATA_FOUND", "NO_DATA_CONFIRMED", "NO_FACILITY_MATCH_CONFIRMED"},
}
ATTACHMENT_SUPPORTED = {".pdf", ".xlsx", ".xls", ".png", ".jpg", ".jpeg"}


def _read_json(path, default=None):
    if not Path(path).is_file():
        return default
    return json.loads(Path(path).read_text(encoding="utf-8-sig"))


def _sha(path):
    sha = hashlib.sha256()
    with Path(path).open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            sha.update(chunk)
    return sha.hexdigest()


def _safe(name):
    value = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", str(name or ""))
    return (value.strip(" .") or "자료")[:115]


def _flatten(row):
    out = {}
    for k, value in row.items():
        if isinstance(value, (list, dict)):
            if k == "cells" and isinstance(value, list):
                for i, cell in enumerate(value, 1):
                    out[f"원문값_{i}"] = cell if not isinstance(cell, (list, dict)) else json.dumps(cell, ensure_ascii=False)
            else:
                out[k] = json.dumps(value, ensure_ascii=False)
        else:
            out[k] = value if value is not None else ""
    return out


def _rows(root, source):
    p = root / source
    if source == "ENVINFO":
        return [("공개연도_사업장", read_csv(p / "discovery.csv")),
                ("첨부자료_원본목록", read_csv(p / "attachment_index.csv"))]
    if source in {"PRTR", "CHEM_STATS"}:
        return [("사업장_연도별", read_csv(p / "discovery.csv")),
                ("원문_상세표", [_flatten(x) for x in read_jsonl(p / "detail_table_rows.jsonl")])]
    if source == "CLEANSYS_AIR":
        return [("연간_원문", [_flatten(x) for x in read_jsonl(p / "annual_rows.jsonl")]),
                ("원문_시설후보", [_flatten(x) for x in (_read_json(p / "candidates.json", []) or [])])]
    if source == "SOOSIRO_WATER":
        return [("연간_원문", [_flatten(x) for x in read_jsonl(p / "annual_rows.jsonl")]),
                ("일별_원문", [_flatten(x) for x in read_jsonl(p / "daily_rows.jsonl")]),
                ("원문_시설후보", [_flatten(x) for x in (_read_json(p / "fact_candidates.json", []) or [])])]
    raise ValueError("Unknown source " + source)


def _valid_pdf(path):
    try:
        reader = PdfReader(path, strict=False)
        return not reader.is_encrypted and len(reader.pages) > 0
    except Exception:
        return False


def _valid_xlsx(path):
    try:
        with zipfile.ZipFile(path) as z:
            if z.testzip():
                return False
        wb = load_workbook(path, read_only=True)
        try:
            return len(wb.sheetnames) >= 1
        finally:
            wb.close()
    except Exception:
        return False


def _copy_attachment(source_path, target, qa):
    ext = source_path.suffix.lower()
    target.parent.mkdir(parents=True, exist_ok=True)
    if ext == ".pdf":
        if not _valid_pdf(source_path):
            qa["problems"].append("ENVINFO_ATTACHMENT_INVALID_PDF:" + str(source_path))
            return False
        shutil.copy2(source_path, target.with_suffix(".pdf"))
        return True
    if ext == ".xlsx":
        if not _valid_xlsx(source_path):
            qa["problems"].append("ENVINFO_ATTACHMENT_INVALID_XLSX:" + str(source_path))
            return False
        shutil.copy2(source_path, target.with_suffix(".xlsx"))
        return True
    if ext in {".png", ".jpg", ".jpeg"}:
        try:
            from PIL import Image
            with Image.open(source_path) as im:
                if im.width < 1 or im.height < 1:
                    raise ValueError("empty image")
                rgb = im.convert("RGB")
                rgb.save(target.with_suffix(".pdf"), "PDF", resolution=120.0)
            if not _valid_pdf(target.with_suffix(".pdf")):
                raise ValueError("converted PDF invalid")
            return True
        except Exception as exc:
            qa["problems"].append("ENVINFO_IMAGE_PDF_CONVERSION_FAILED:" + str(source_path) + ":" + str(exc))
            return False
    qa["problems"].append("ENVINFO_UNSUPPORTED_ATTACHMENT_TYPE:" + str(source_path))
    return False


def _render_envinfo(root, target, qa):
    env = root / "ENVINFO"
    discovered = read_csv(env / "discovery.csv")
    expected_keys = set()
    for entry in discovered:
        year, site = str(entry.get("year") or "").strip(), str(entry.get("compId") or "").strip()
        if year and site:
            expected_keys.add((year, site))
    files = sorted((env / "raw_detail").glob("*.html"))
    rendered = set()
    for year, site in sorted(expected_keys):
        found = [p for p in files if p.name.startswith(f"{year}_{site}_")]
        if not found:
            qa["problems"].append(f"ENVINFO_MISSING_HTML:{year}:{site}")
            continue
        p = target / "ENVINFO_공개정보본문" / f"{year}_{site}_공개정보본문_인쇄본.pdf"
        p.parent.mkdir(parents=True, exist_ok=True)
        ok, error = render_envinfo_static_pdf(found[0], p)
        if not ok or not _valid_pdf(p):
            p.unlink(missing_ok=True)
            qa["problems"].append(f"ENVINFO_BODY_PDF_FAILED:{year}:{site}:{error}")
            continue
        rendered.add((year, site))
    qa["envinfo_disclosure_expected"] = len(expected_keys)
    qa["envinfo_disclosure_pdf_created"] = len(rendered)
    if expected_keys != rendered:
        qa["problems"].append("ENVINFO_DISCLOSURE_PDF_NOT_COMPLETE")
    return len(rendered)


def _attachments(root, target, qa):
    env = root / "ENVINFO"
    indexed_rows = read_csv(env / "attachment_index.csv")
    rows = [r for r in indexed_rows if r.get("collection_status") == "DOWNLOADED"]
    missing = [r for r in indexed_rows if r.get("collection_status") != "DOWNLOADED"]
    if missing:
        qa["problems"].append(f"ENVINFO_ATTACHMENT_DOWNLOAD_NOT_COMPLETE:{len(missing)}")
    # One source-native record must result in one user-visible attachment,
    # even when the collector stores identical bytes only once across years.
    # S-OIL's 2024 policy attachment refers to the 2023 canonical PDF file.
    exported = 0
    unique_origins = set()
    for row_idx, row in enumerate(rows, 1):
        source = root.parent / str(row.get("stored_path") or "")
        # Source-native collector indices normally contain `output/ENVINFO/...`.
        if not source.exists():
            source = root.parent.parent / str(row.get("stored_path") or "")
        original = str(row.get("original_filename") or source.name)
        year = str(row.get("year") or "연도미상")
        site = str(row.get("compId") or "사업장미상")
        unique_origins.add((str(source), str(row.get("sha256") or "")))
        if not source.is_file():
            qa["problems"].append("ENVINFO_ATTACHMENT_MISSING:" + str(source))
            continue
        digest = str(row.get("sha256") or "").strip()
        if digest and digest != _sha(source):
            qa["problems"].append("ENVINFO_ATTACHMENT_SHA_MISMATCH:" + str(source))
            continue
        ext = source.suffix.lower()
        if ext == ".xls":
            qa["problems"].append("ENVINFO_LEGACY_XLS_REQUIRES_CONVERSION:" + str(source))
            continue
        target_file = (target / "ENVINFO_공식첨부자료" / _safe(year) / _safe(site) /
                       f"{row_idx:03d}_{_safe(Path(original).stem)}")
        if _copy_attachment(source, target_file, qa):
            exported += 1
    qa["envinfo_source_attachments_downloaded"] = len(rows)
    qa["envinfo_attachment_unique_origins"] = len(unique_origins)
    qa["envinfo_attachments_delivered"] = exported
    if exported != len(rows):
        qa["problems"].append("ENVINFO_ATTACHMENT_DELIVERY_NOT_COMPLETE")


def _reports(archive, target, qa, required_report_years=()):
    if not archive:
        qa["report_status"] = "NOT_PROVIDED"
        qa["limitations"].append("CORPORATE_REPORTS_NOT_INCLUDED_IN_PUBLIC_SOURCE_ONLY_PACKAGE")
        return
    archive = Path(archive)
    if not archive.is_file():
        qa["problems"].append("REPORT_PDF_ARCHIVE_MISSING")
        return
    count = 0
    years_delivered = set()
    try:
        with zipfile.ZipFile(archive) as zipped:
            if zipped.testzip():
                raise ValueError("report ZIP integrity error")
            for entry in zipped.infolist():
                if entry.is_dir():
                    continue
                if Path(entry.filename).suffix.lower() != ".pdf":
                    qa["problems"].append("NON_PDF_INSIDE_REPORT_ARCHIVE:" + entry.filename)
                    continue
                path = target / "공식_지속가능경영보고서_PDF" / _safe(Path(entry.filename).name)
                path.parent.mkdir(parents=True, exist_ok=True)
                with path.open("wb") as f, zipped.open(entry) as source:
                    shutil.copyfileobj(source, f)
                if not _valid_pdf(path):
                    qa["problems"].append("INVALID_OFFICIAL_REPORT_PDF:" + entry.filename)
                    path.unlink(missing_ok=True)
                else:
                    count += 1
                    years_delivered.update(re.findall(r"(?<!\d)20\d{2}(?!\d)", Path(entry.filename).name))
        qa["official_report_pdfs_delivered"] = count
        qa["official_report_years_delivered"] = sorted(years_delivered)
        missing_years = sorted(set(map(str, required_report_years)) - years_delivered)
        if missing_years:
            qa["problems"].append("OFFICIAL_ANNUAL_REPORT_YEARS_MISSING:" + ",".join(missing_years))
        if count == 0:
            qa["problems"].append("NO_OFFICIAL_REPORT_PDFS")
    except (OSError, zipfile.BadZipFile, ValueError) as exc:
        qa["problems"].append("REPORT_ZIP_UNREADABLE:" + str(exc))


def create(root, destination, company, reports_archive=None, required_report_years=()):
    root = Path(root).resolve()
    dest = Path(destination).resolve()
    dest.mkdir(parents=True, exist_ok=True)
    staging = dest / "pdf_excel_files"
    if staging.exists():
        shutil.rmtree(staging)
    staging.mkdir()
    qa = {"company": company, "status": "DELIVERY_NOT_READY",
          "problems": [], "limitations": [], "source_checks": {}, "required_sources": list(SOURCES),
          "outputs": {}}
    overview = []
    for source in SOURCES:
        src = root / source
        status = _read_json(src / "status.json", {}) or {}
        state = str(status.get("status") or "MISSING_STATUS")
        source_rows = _rows(root, source)
        counts = {name: len(rows) for name, rows in source_rows}
        errors = status.get("errors", [])
        has_errors = (bool(errors) if isinstance(errors, (list, dict, str))
                      else int(errors or 0) > 0)
        acceptable = state in ACCEPTED_SOURCE_STATUSES[source] and src.is_dir() and not has_errors
        snapshot = _read_json(src / "source_snapshot_recovery.json", {}) or {}
        if snapshot:
            from orchestrator.source_snapshot_recovery import checksum
            if (snapshot.get("recovery_status") != "VERIFIED_ORIGINAL_SNAPSHOT_REUSED"
                    or snapshot.get("source") != source
                    or not snapshot.get("source_run_id")
                    or not snapshot.get("source_captured_at")):
                qa["problems"].append("UNVERIFIED_SNAPSHOT_RECOVERY:" + source)
            else:
                digests = snapshot.get("original_files_sha256") or {}
                if source != "PRTR" or not isinstance(digests, dict) or len(digests) < 3:
                    qa["problems"].append("INCOMPLETE_OR_UNSUPPORTED_SNAPSHOT_RECEIPT:" + source)
                else:
                    for relative, expected_digest in digests.items():
                        source_rel = Path(relative)
                        if source_rel.is_absolute() or ".." in source_rel.parts or not str(relative):
                            qa["problems"].append("UNSAFE_SNAPSHOT_PROOF_PATH:" + str(relative))
                            continue
                        original = src / source_rel
                        if not original.is_file() or checksum(original) != expected_digest:
                            qa["problems"].append("SNAPSHOT_ORIGINAL_CHANGED:" + source + ":" + relative)
                qa["limitations"].append(
                    source + ":VERIFIED_PRIOR_SNAPSHOT_USED_NOT_CURRENT_COLLECTION"
                )
        if state in {"NO_FACILITY_MATCH_CONFIRMED", "NO_PUBLIC_RECORD_FOUND_IN_SCOPE", "NO_DATA_CONFIRMED"}:
            qa["limitations"].append(source + ":NO_SOURCE_NATIVE_ROWS_OR_FACILITY_MATCH")
        if state == "DATA_FOUND" and not sum(counts.values()):
            qa["problems"].append("DATA_FOUND_WITHOUT_USER_SOURCE_ROWS:" + source)
        for flag in ("detail_fail", "attachment_fail"):
            if int(status.get(flag, 0) or 0) > 0:
                qa["problems"].append(f"COLLECTOR_PARTIAL_DETAIL_FAILURE:{source}:{flag}")
        qa["source_checks"][source] = {
            "collector_status": state, "counts": counts,
            "evidence_present": src.is_dir(),
            "source_url": SOURCE_URLS[source],
            "retrieval_status": ("검증된 과거 원본 재사용" if snapshot else "이번 실행 수집본"),
            "original_capture_date": snapshot.get("source_captured_at", "") if snapshot else "",
            "original_run_id": snapshot.get("source_run_id", "") if snapshot else "",
        }
        if not acceptable:
            qa["problems"].append("SOURCE_INCOMPLETE:" + source + ":" + state)
        overview.append({
            "자료원": source, "수집상태": state,
            "원자료_합계행수": sum(counts.values()),
            "상세내역": "; ".join(f"{k} {v}행" for k, v in counts.items()),
            "공식_출처_URL": SOURCE_URLS[source],
            "자료취득경로": ("과거 검증본 재사용 (현재 서버 재조회 실패)" if snapshot else "이번 실행에서 원본 수집"),
            "과거원본수집일시": snapshot.get("source_captured_at", "") if snapshot else "",
            "과거원본GitHub실행ID": snapshot.get("source_run_id", "") if snapshot else "",
            "주의": ("자료원별 식별자는 서로 다른 체계임. 같은 주소만으로 동일 시설이라고 확정하지 않음."
                   if source != "ENVINFO" else "본문 PDF는 원본 HTML을 재구성한 인쇄본. 공식 원본 PDF와 구별.")
        })
        try:
            dict_rows_to_xlsx(staging / f"{source}_원문_표.xlsx",
                              [("수집상태", [overview[-1]])] + source_rows)
        except Exception as exc:
            qa["problems"].append("EXCEL_EXPORT_FAILED:" + source + ":" + str(exc))
    # Human interpretation layer is separate from raw exports: it does not
    # rewrite source measurements, promote unresolved PRTR identity or assert
    # comparability from an unverified TMS unit.
    interpretation = review_environment_evidence(root)
    qa["interpretation"] = {
        "status": interpretation["status"],
        "flags": interpretation["flags"],
        "qa_issues": interpretation["qa_issues"],
        "identity_records": len(interpretation["identity_rows"]),
        "legal_disclosure_records": len(interpretation["violation_rows"]),
        "legal_disclosure_source_pages": len(interpretation["legal_coverage_rows"]),
        "legal_pages_without_individual_records": sum(
            not x["개별위반기록추출수"] for x in interpretation["legal_coverage_rows"]),
        "prtr_years_compared": len(interpretation["prtr_comparison_rows"]),
        "air_years_reviewed": len(interpretation["air_comparison_rows"]),
    }
    for flag in interpretation["flags"] + interpretation["qa_issues"]:
        qa["limitations"].append("EVIDENCE_REVIEW:" + flag)
    dict_rows_to_xlsx(
        staging / "01_환경규제_사업장식별_수치검토.xlsx",
        [("검토요약", interpretation["overview"]),
         ("사업장_식별근거", interpretation["identity_rows"]),
         ("행정처분_원문", interpretation["violation_rows"]),
         ("행정처분_공개범위", interpretation["legal_coverage_rows"]),
         ("PRTR_화학배출량비교", interpretation["prtr_comparison_rows"]),
         ("대기TMS_단위범위검토", interpretation["air_comparison_rows"]),
         ("ENVINFO_수치원문", interpretation["envinfo_metric_rows"])],
    )
    prtr_unresolved = any(
        r["자료원"] == "PRTR" and
        any(status != "MATCHED" for status in r["식별상태"].split(" / "))
        for r in interpretation["identity_rows"]
    )
    for row in overview:
        if row["자료원"] == "PRTR" and prtr_unresolved:
            row["주의"] = (
                "PRTR 식별 미확정: UNRESOLVED 또는 주소 후보(ADDRESS_CANDIDATE) "
                "상태는 공식 사업장 고유 ID 연계를 입증하지 않음"
            )
        if row["자료원"] == "PRTR" and qa["source_checks"]["PRTR"]["retrieval_status"] == "검증된 과거 원본 재사용":
            row["주의"] += "; 이번 PRTR 재조회 실패, 기존 검증 자료와 수집시점 표시"
        if row["자료원"] == "CLEANSYS_AIR" and interpretation["air_comparison_rows"]:
            row["주의"] = "수치 필드 단위 및 ENVINFO와 집계범위 동일성 미검증. 배출량 차이 임의 계산 금지"
    dict_rows_to_xlsx(staging / "00_수집범위_및_자료목록.xlsx",
                      [("자료원별_수집현황", overview),
                       ("중요_검토사항", interpretation["overview"]),
                       ("원본_사업장식별", interpretation["identity_rows"]),
                       ("자료원_구분_주의", [
                         {"확인사항": "공식 공개 여부", "설명": "공개 목록 미매칭은 TMS 비대상 또는 배출량 0의 증거가 아닙니다."},
                         {"확인사항": "자료원 원본", "설명": "공식 사이트 자료를 보존·표로 정리했으며 수치 단위를 임의로 환산하지 않았습니다."},
                         {"확인사항": "완료 판단", "설명": "수집 절차 성공과 원문 확보 및 PDF/Excel 전달 완료는 별도 검사합니다."},
                       ])])
    _render_envinfo(root, staging, qa)
    _attachments(root, staging, qa)
    _reports(reports_archive, staging, qa, required_report_years)
    files = sorted(p for p in staging.rglob("*") if p.is_file())
    counts = Counter(p.suffix.lower() for p in files)
    if not files or counts[".xlsx"] < 7:
        qa["problems"].append("REQUIRED_SEVEN_EXCEL_WORKBOOKS_MISSING")
    if any(p.suffix.lower() not in USER_EXTENSIONS for p in files):
        qa["problems"].append("NON_PDF_EXCEL_USER_DELIVERY_FILE")
    for p in files:
        if p.suffix.lower() == ".pdf" and not _valid_pdf(p):
            qa["problems"].append("DELIVERED_PDF_UNREADABLE:" + str(p.relative_to(staging)))
        if p.suffix.lower() == ".xlsx" and not _valid_xlsx(p):
            qa["problems"].append("DELIVERED_EXCEL_UNREADABLE:" + str(p.relative_to(staging)))
    zip_path = dest / "Public_Environmental_PDF_Excel.zip"
    with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_STORED) as z:
        for p in files:
            z.write(p, p.relative_to(staging).as_posix())
    with zipfile.ZipFile(zip_path) as z:
        if z.testzip() or len([n for n in z.namelist() if not n.endswith("/")]) != len(files):
            qa["problems"].append("ZIP_INTEGRITY_OR_FILE_COUNT_ERROR")
    qa["outputs"] = {"zip": str(zip_path), "zip_sha256": _sha(zip_path),
                     "zip_bytes": zip_path.stat().st_size,
                     "pdf_files": counts[".pdf"], "excel_files": counts[".xlsx"]}
    qa["status"] = ("DELIVERY_REVIEW_REQUIRED" if qa["problems"] else
                    "USER_PDF_EXCEL_ZIP_READY_WITH_LIMITATIONS" if qa["limitations"] else
                    "USER_PDF_EXCEL_ZIP_READY")
    report = dest / "Internal_Delivery_QA.json"
    report.write_text(json.dumps(qa, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"delivery_status": qa["status"],
                      "file_counts": qa["outputs"],
                      "problem_count": len(qa["problems"]),
                      "problems": qa["problems"][:15]}, ensure_ascii=False), flush=True)
    return qa


def main():
    p = argparse.ArgumentParser()
    p.add_argument("--output-root", required=True, help="Parent containing ENVINFO, PRTR, CHEM_STATS, CLEANSYS_AIR, SOOSIRO_WATER")
    p.add_argument("--out", required=True)
    p.add_argument("--company", required=True)
    p.add_argument("--reports-zip", default="")
    p.add_argument("--required-report-years", default="",
                   help="Comma-separated official annual report years that must be present")
    args = p.parse_args()
    years = tuple(y.strip() for y in args.required_report_years.split(",") if y.strip())
    result = create(args.output_root, args.out, args.company, args.reports_zip or None, years)
    return 0 if result["status"] in {
        "USER_PDF_EXCEL_ZIP_READY", "USER_PDF_EXCEL_ZIP_READY_WITH_LIMITATIONS"
    } else 2


if __name__ == "__main__":
    raise SystemExit(main())
