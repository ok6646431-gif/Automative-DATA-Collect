"""Build a human-readable review package even when strict package validation fails.

This is a delivery fallback, not a validation bypass. The strict package/Archive gates
remain authoritative. The review package exposes only readable summaries and verified
company documents; machine/raw evidence stays in the normal final package.
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import shutil
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
import sys
sys.path.insert(0, str(ROOT / "orchestrator"))
from archive_builder import render_html_pdf, render_url_pdf, valid_pdf  # noqa: E402

SOURCE_LABELS = {
    "ENVINFO": "환경정보공개시스템",
    "PRTR": "화학물질 배출량조사(PRTR)",
    "CHEM_STATS": "화학물질 통계조사",
    "CLEANSYS_AIR": "CleanSYS 대기TMS",
    "SOOSIRO_WATER": "수질TMS",
}


def read_json(path: Path, default=None):
    if not path.exists():
        return default
    return json.loads(path.read_text(encoding="utf-8"))


def read_csv(path: Path):
    if not path.exists() or path.stat().st_size == 0:
        return []
    with path.open(encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def read_jsonl(path: Path):
    if not path.exists() or path.stat().st_size == 0:
        return []
    rows = []
    for line in path.read_text(encoding="utf-8", errors="replace").splitlines():
        if not line.strip():
            continue
        try:
            value = json.loads(line)
        except json.JSONDecodeError:
            continue
        if isinstance(value, dict):
            rows.append(value)
    return rows


def safe_name(value: str) -> str:
    text = re.sub(r'[\\/:*?"<>|]+', "_", str(value or "")).strip()
    return text[:120] or "자료"


def style_sheet(ws, freeze="A2"):
    if ws.max_row:
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = freeze
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for idx, col in enumerate(ws.columns, 1):
        values = [str(c.value or "") for c in col[:80]]
        width = min(42, max(10, max((len(v) for v in values), default=8) + 2))
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.auto_filter.ref = ws.dimensions


def excel_value(value):
    if value is None:
        return ""
    if isinstance(value, (dict, list, tuple, set)):
        return json.dumps(value, ensure_ascii=False, sort_keys=isinstance(value, dict))
    return value


def write_rows_sheet(wb: Workbook, title: str, rows: list[dict], preferred=None):
    ws = wb.create_sheet(title[:31])
    if not rows:
        ws.append(["상태", "내용"])
        ws.append(["NO_DATA", "확인된 행이 없습니다."])
        style_sheet(ws)
        return ws
    fields = []
    for key in preferred or []:
        if any(key in row for row in rows) and key not in fields:
            fields.append(key)
    for row in rows:
        for key in row:
            if key not in fields:
                fields.append(key)
    ws.append(fields)
    for row in rows:
        ws.append([excel_value(row.get(k, "")) for k in fields])
    style_sheet(ws)
    return ws


def source_status_rows(assembled: Path):
    manifest = read_json(assembled / "Master_Manifest.json", {}) or {}
    rows = []
    for source, label in SOURCE_LABELS.items():
        status = read_json(assembled / "output" / source / "status.json", {}) or {}
        rows.append({
            "자료원": label,
            "source_key": source,
            "상태": status.get("status", "MISSING_STATUS"),
            "조회요청수": status.get("requests", ""),
            "발견행수": status.get("rows", status.get("annual_rows", "")),
            "상세확보": status.get("detail_ok", ""),
            "오류": status.get("errors", status.get("detail_fail", "")),
            "비고": " | ".join(str(x) for x in (manifest.get("sources", {}).get(source, {}).get("checks") or [])),
        })
    corp = read_json(assembled / "output" / "CORP_DOCS" / "status.json", {}) or {}
    rows.append({
        "자료원": "회사 공식문서",
        "source_key": "CORP_DOCS",
        "상태": corp.get("status", "MISSING_STATUS"),
        "조회요청수": corp.get("documents_declared", ""),
        "발견행수": corp.get("downloaded", ""),
        "상세확보": corp.get("downloaded", ""),
        "오류": corp.get("failed", ""),
        "비고": "",
    })
    return rows


def review_rows(assembled: Path):
    rows = []
    for item in read_csv(assembled / "Validation_Queue.csv"):
        if str(item.get("status") or "").upper() != "REVIEW_REQUIRED":
            continue
        rows.append({
            "구분": item.get("object_type", ""),
            "대상": item.get("object_key", ""),
            "문제유형": item.get("issue_type", ""),
            "중요도": item.get("severity", ""),
            "근거": item.get("evidence", ""),
            "확인/조치": item.get("recommended_action", ""),
            "비고": item.get("notes", ""),
        })
    seen = {(r["구분"], r["대상"], r["문제유형"]) for r in rows}
    for item in read_json(assembled / "REVIEW_REQUIRED.json", []) or []:
        if not isinstance(item, dict):
            continue
        row = {
            "구분": item.get("object_type", item.get("source", "SOURCE")),
            "대상": item.get("object_key", item.get("source", "")),
            "문제유형": item.get("issue_type", " | ".join(item.get("issues") or [])),
            "중요도": item.get("severity", ""),
            "근거": item.get("evidence", item.get("status", "")),
            "확인/조치": item.get("recommended_action", "원자료와 식별 근거를 확인"),
            "비고": item.get("notes", ""),
        }
        key = (row["구분"], row["대상"], row["문제유형"])
        if key not in seen:
            rows.append(row)
            seen.add(key)
    return rows


def make_status_workbook(assembled: Path, out: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "수집현황"
    manifest = read_json(assembled / "Master_Manifest.json", {}) or {}
    profile = read_json(assembled / "Company_Profile.json", {}) or {}
    state = "AUTO_COMPLETE"
    if manifest.get("package_health") == "FAIL":
        state = "BLOCKED"
    elif manifest.get("validation") == "REVIEW_REQUIRED" or review_rows(assembled):
        state = "REVIEW_REQUIRED"
    summary = [
        ("회사", profile.get("company_display_name", profile.get("requested_company_name", ""))),
        ("최종상태", state),
        ("package_health", manifest.get("package_health", "UNKNOWN")),
        ("validation", manifest.get("validation", "UNKNOWN")),
        ("확인필요 건수", len(review_rows(assembled))),
        ("안내", "BLOCKED/REVIEW_REQUIRED라도 아래 자료는 검증 상태를 구분해 제공하며, 시스템 raw는 이 패키지에 포함하지 않습니다."),
    ]
    ws.append(["항목", "값"])
    for key, value in summary:
        ws.append([key, value])
    style_sheet(ws)
    write_rows_sheet(
        wb, "자료원별 상태", source_status_rows(assembled),
        ["자료원", "상태", "발견행수", "상세확보", "오류", "비고"],
    )
    docs = read_csv(assembled / "output" / "CORP_DOCS" / "document_index.csv")
    write_rows_sheet(
        wb, "회사공식자료", docs,
        ["title", "document_type", "report_year", "verification_status", "collection_status", "source_url"],
    )
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return state


def make_review_workbook(assembled: Path, out: Path):
    wb = Workbook()
    wb.remove(wb.active)
    rows = review_rows(assembled)
    write_rows_sheet(
        wb, "확인필요", rows,
        ["구분", "대상", "문제유형", "중요도", "근거", "확인/조치", "비고"],
    )
    identities = read_csv(assembled / "Source_Identity.csv")
    candidates = [r for r in identities if str(r.get("match_status") or "").upper() != "CONFIRMED"]
    write_rows_sheet(
        wb, "미확정_사업장후보", candidates,
        ["source_key", "source_site_id", "source_site_name_raw", "source_address_raw",
         "match_status", "match_method", "match_score"],
    )
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)


def _load_source_rows(root: Path, source: str):
    base = root / "output" / source
    rows = read_csv(base / "discovery.csv")
    if rows:
        return rows
    if source == "CLEANSYS_AIR":
        return read_jsonl(base / "annual_rows.jsonl") or (read_json(base / "candidates.json", []) or [])
    if source == "SOOSIRO_WATER":
        return read_jsonl(base / "annual_rows.jsonl") or (read_json(base / "fact_candidates.json", []) or [])
    return []


def make_public_workbook(assembled: Path, out: Path):
    wb = Workbook()
    wb.remove(wb.active)
    write_rows_sheet(
        wb, "자료원_상태", source_status_rows(assembled),
        ["자료원", "상태", "발견행수", "상세확보", "오류", "비고"],
    )
    for source in SOURCE_LABELS:
        rows = _load_source_rows(assembled, source)
        if not rows:
            status = read_json(assembled / "output" / source / "status.json", {}) or {}
            rows = [{
                "상태": status.get("status", "NO_DATA"),
                "설명": "확정된 공개자료 행이 없습니다. 수집 상태는 자료원_상태 시트를 확인하십시오.",
            }]
        write_rows_sheet(wb, source[:31], rows)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)


def copy_company_documents(assembled: Path, target: Path):
    target.mkdir(parents=True, exist_ok=True)
    rows = read_csv(assembled / "output" / "CORP_DOCS" / "document_index.csv")
    delivered = []
    for row in rows:
        if str(row.get("collection_status") or "").upper() != "DOWNLOADED":
            continue
        if str(row.get("verification_status") or "").upper() not in {"VERIFIED", "SOURCE_VERIFIED"}:
            continue
        title = safe_name(row.get("title") or row.get("document_id") or "회사자료")
        year = str(row.get("report_year") or "").strip()
        stem = safe_name((year + "_" if year else "") + title)
        stored = assembled / str(row.get("stored_path") or "")
        out_path = target / f"{stem}.pdf"
        note = ""
        if stored.exists() and stored.suffix.casefold() == ".pdf" and valid_pdf(stored):
            shutil.copy2(stored, out_path)
        else:
            ok = False
            url = str(row.get("source_url") or row.get("source_locator") or "").strip()
            if url.startswith("http"):
                ok, note = render_url_pdf(url, out_path)
            if not ok and stored.exists():
                ok, note2 = render_html_pdf(stored, out_path)
                note = note or note2
            if not ok:
                txt = target / f"{stem}_원문링크.txt"
                txt.write_text(
                    f"{row.get('title', '')}\n\n원문: {url}\n\nPDF 변환 실패: {note}\n",
                    encoding="utf-8",
                )
                out_path = txt
        delivered.append({
            "title": row.get("title", ""),
            "report_year": year,
            "delivered_path": out_path.name,
            "source_url": row.get("source_url", ""),
            "verification_status": row.get("verification_status", ""),
        })
    return delivered


def build(assembled: Path, out_dir: Path, prefix: str):
    assembled = assembled.resolve()
    out_dir = out_dir.resolve()
    if out_dir.exists():
        shutil.rmtree(out_dir)
    package = out_dir / safe_name(prefix)
    first = package / "00_먼저보기"
    corp = package / "01_회사공식자료"
    public = package / "02_공공환경자료"
    first.mkdir(parents=True, exist_ok=True)
    public.mkdir(parents=True, exist_ok=True)

    state = make_status_workbook(assembled, first / "수집현황.xlsx")
    make_review_workbook(assembled, first / "확인필요_REVIEW_REQUIRED.xlsx")
    make_public_workbook(assembled, public / "공공환경자료_정리.xlsx")
    docs = copy_company_documents(assembled, corp)

    (first / "README_먼저읽기.txt").write_text(
        "이 패키지는 사람용 검토 자료입니다.\n"
        "- AUTO_COMPLETE: 엄격 검증 통과\n"
        "- REVIEW_REQUIRED: 자료는 제공하되 일부 식별/범위 확인 필요\n"
        "- BLOCKED: 엄격 패키지 검증 실패. 확인된 자료만 참고하고 확인필요.xlsx를 먼저 검토\n\n"
        "JSON, raw HTML, 실행로그 등 시스템 원본은 이 패키지에서 제외했습니다.\n",
        encoding="utf-8",
    )

    zip_path = Path(
        shutil.make_archive(
            str(out_dir / safe_name(prefix)),
            "zip",
            root_dir=out_dir,
            base_dir=package.name,
        )
    )
    result = {
        "schema_version": "human-review-package-1.0",
        "status": "PASS",
        "review_state": state,
        "package_root": package.name,
        "zip_path": zip_path.name,
        "company_documents": len(docs),
        "review_items": len(review_rows(assembled)),
        "principle": "Strict validation state is preserved; this package only makes verified/review-required evidence readable.",
    }
    (out_dir / "review_package_result.json").write_text(
        json.dumps(result, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return result


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--assembled", default="assembled", type=Path)
    ap.add_argument("--out-dir", default="review-delivery", type=Path)
    ap.add_argument("--prefix", default="환경자료_검토용")
    args = ap.parse_args()
    print(json.dumps(build(args.assembled, args.out_dir, args.prefix), ensure_ascii=False))


if __name__ == "__main__":
    main()
