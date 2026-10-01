import csv
import json
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "tools"))
import build_human_review_package as review_pkg


def write_json(path, payload):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")


def write_csv(path, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    fields = []
    for row in rows:
        for key in row:
            if key not in fields:
                fields.append(key)
    if not fields:
        fields = ["value"]
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        w.writerows(rows)


class HumanReviewPackageTests(unittest.TestCase):
    def test_blocked_run_still_yields_readable_review_package_without_raw_json(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            assembled = root / "assembled"
            assembled.mkdir()
            write_json(assembled / "Master_Manifest.json", {
                "package_health": "FAIL",
                "validation": "REVIEW_REQUIRED",
                "sources": {"ENVINFO": {"checks": ["zero_byte_artifact"]}},
            })
            write_json(assembled / "Company_Profile.json", {
                "company_display_name": "테스트건설",
                "requested_company_name": "테스트건설",
            })
            write_json(assembled / "REVIEW_REQUIRED.json", [{
                "source": "ENVINFO",
                "issues": ["zero_byte_artifact"],
                "status": "NO_MATCH",
            }])
            write_csv(assembled / "Validation_Queue.csv", [{
                "validation_id": "V1",
                "object_type": "SOURCE_IDENTITY",
                "object_key": "CHEM_STATS:X1",
                "issue_type": "IDENTITY_REVIEW",
                "severity": "HIGH",
                "evidence": "동명이인 후보",
                "recommended_action": "법인 식별 확인",
                "status": "REVIEW_REQUIRED",
                "notes": "",
            }])
            write_csv(assembled / "Source_Identity.csv", [{
                "source_key": "CHEM_STATS",
                "source_site_id": "X1",
                "source_site_name_raw": "테스트산업",
                "source_address_raw": "부산",
                "match_status": "REVIEW_REQUIRED",
            }])

            for source, status in {
                "ENVINFO": "NO_MATCH",
                "PRTR": "NO_MATCH",
                "CHEM_STATS": "DATA_FOUND",
                "CLEANSYS_AIR": "NO_FACILITY_MATCH_CONFIRMED",
                "SOOSIRO_WATER": "NO_FACILITY_MATCH_CONFIRMED",
            }.items():
                payload = {"status": status, "rows": 0}
                if source == "CLEANSYS_AIR":
                    payload["errors"] = []
                write_json(assembled / "output" / source / "status.json", payload)
            write_csv(assembled / "output" / "CHEM_STATS" / "discovery.csv", [{
                "search_year": "2024",
                "bplcId": "X1",
                "bplcNm": "테스트산업",
            }])

            html = assembled / "output" / "CORP_DOCS" / "raw_documents" / "ENVIRONMENTAL_POLICY" / "2026" / "policy.asp"
            html.parent.mkdir(parents=True, exist_ok=True)
            html.write_text("<html><body>환경경영방침</body></html>", encoding="utf-8")
            write_json(assembled / "output" / "CORP_DOCS" / "status.json", {
                "status": "DATA_FOUND", "documents_declared": 1, "downloaded": 1, "failed": 0,
            })
            write_csv(assembled / "output" / "CORP_DOCS" / "document_index.csv", [{
                "document_id": "D1",
                "document_type": "ENVIRONMENTAL_POLICY",
                "title": "환경경영방침",
                "report_year": "2026",
                "stored_path": str(html.relative_to(assembled)),
                "source_url": "https://official.example/policy",
                "verification_status": "SOURCE_VERIFIED",
                "collection_status": "DOWNLOADED",
            }])

            out = root / "review"
            with patch.object(review_pkg, "render_url_pdf", return_value=(False, "offline")), \
                 patch.object(review_pkg, "render_html_pdf", return_value=(False, "no browser")):
                result = review_pkg.build(assembled, out, "테스트건설_검토용")

            self.assertEqual(result["status"], "PASS")
            self.assertEqual(result["review_state"], "BLOCKED")
            zip_path = out / result["zip_path"]
            self.assertTrue(zip_path.exists())

            with zipfile.ZipFile(zip_path) as zf:
                names = [n for n in zf.namelist() if not n.endswith("/")]
                self.assertTrue(any(n.endswith("00_먼저보기/수집현황.xlsx") for n in names))
                self.assertTrue(any(n.endswith("00_먼저보기/확인필요_REVIEW_REQUIRED.xlsx") for n in names))
                self.assertTrue(any(n.endswith("02_공공환경자료/공공환경자료_정리.xlsx") for n in names))
                self.assertTrue(any(n.endswith("01_회사공식자료/2026_환경경영방침_원문링크.txt") for n in names))
                self.assertFalse(any(n.lower().endswith((".json", ".jsonl", ".html", ".htm", ".do")) for n in names))

            status_book = load_workbook(out / "테스트건설_검토용" / "00_먼저보기" / "수집현황.xlsx")
            values = {status_book["수집현황"][f"A{i}"].value: status_book["수집현황"][f"B{i}"].value for i in range(2, 8)}
            self.assertEqual(values["최종상태"], "BLOCKED")

            review_book = load_workbook(out / "테스트건설_검토용" / "00_먼저보기" / "확인필요_REVIEW_REQUIRED.xlsx")
            self.assertGreaterEqual(review_book["확인필요"].max_row, 2)

            public_book = load_workbook(out / "테스트건설_검토용" / "02_공공환경자료" / "공공환경자료_정리.xlsx")
            confirmed_values = [
                cell.value
                for row in public_book["화학통계_확정"].iter_rows()
                for cell in row
            ]
            review_values = [
                cell.value
                for row in public_book["화학통계_검토필요"].iter_rows()
                for cell in row
            ]
            self.assertNotIn("X1", confirmed_values)
            self.assertIn("X1", review_values)
            self.assertIn("REVIEW_REQUIRED", review_values)


if __name__ == "__main__":
    unittest.main()
