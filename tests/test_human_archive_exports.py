import csv
import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from orchestrator.human_archive_exports import build_human_excels


class HumanArchiveExportsTest(unittest.TestCase):
    def _json(self, path, value):
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(json.dumps(value, ensure_ascii=False), encoding="utf-8")

    def _jsonl(self, path, rows):
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text("\n".join(json.dumps(r, ensure_ascii=False) for r in rows) + ("\n" if rows else ""), encoding="utf-8")

    def _csv(self, path, rows):
        path.parent.mkdir(parents=True, exist_ok=True)
        fields=[]
        for r in rows:
            for k in r:
                if k not in fields: fields.append(k)
        if not fields: fields=["value"]
        with path.open("w",encoding="utf-8-sig",newline="") as f:
            w=csv.DictWriter(f,fieldnames=fields); w.writeheader(); w.writerows(rows)

    def test_unbound_soosiro_is_quarantined_not_silently_dropped(self):
        with tempfile.TemporaryDirectory() as td:
            package=Path(td)/"assembled"; out=package/"output"; archive=package/"Human"
            self._csv(package/"Source_Identity.csv", [{
                "source_key":"SOOSIRO_WATER","source_site_id":"31F0081","source_site_name_raw":"효성티앤씨(주) 울산공장",
                "source_address_raw":"울산광역시 남구 매암동 583-1","valid_from":"2020","valid_to":"2025",
                "match_status":"REVIEW_REQUIRED","match_basis":"SINGLE_OR_CONFLICTING_SOURCE_ADDRESS","review_required":"True","notes":"address/name pair not independently confirmed"
            }])
            annual={"YEAR":"2020","FACT_CODE":"31F0081","FACT_FNAME":"효성티앤씨(주) 울산공장","FACT_ADDR":"울산광역시 남구 매암동 583-1","COD_AVRG_DNSTY":"9.6","COD_DSCAMT":"11984","TN_AVRG_DNSTY":"8.102","TN_DSCAMT":"9935","SS_AVRG_DNSTY":"1.3","SS_DSCAMT":"1334","TP_AVRG_DNSTY":"0.148","TP_DSCAMT":"1","WAST_NO":1}
            daily={**annual,"YEAR":"2024","DAY":"01월01일","QUARTER":"1분기","AMOUNT_FLOW":"2098.7","source_fact_code":"31F0081"}
            self._jsonl(out/"SOOSIRO_WATER"/"annual_rows.jsonl",[annual])
            self._jsonl(out/"SOOSIRO_WATER"/"daily_rows.jsonl",[daily])
            self._json(out/"SOOSIRO_WATER"/"fact_candidates.json",[{"FACT_CODE":"31F0081","FACT_FNAME":"효성티앤씨(주) 울산공장","FACT_ADDR":"울산광역시 남구 매암동 583-1"}])
            self._json(out/"SOOSIRO_WATER"/"status.json",{"status":"DATA_FOUND"})
            for src in ["CLEANSYS_AIR","PRTR","CHEM_STATS"]:
                self._json(out/src/"status.json",{"status":"NO_MATCH"})
            self._jsonl(out/"CLEANSYS_AIR"/"annual_rows.jsonl",[])
            self._json(out/"CLEANSYS_AIR"/"candidates.json",[])
            self._csv(out/"PRTR"/"discovery.csv",[]); self._jsonl(out/"PRTR"/"detail_table_rows.jsonl",[])
            self._csv(out/"CHEM_STATS"/"discovery.csv",[]); self._jsonl(out/"CHEM_STATS"/"detail_table_rows.jsonl",[])

            build_human_excels(package,archive,{s:set() for s in ["CLEANSYS_AIR","SOOSIRO_WATER","PRTR","CHEM_STATS"]})
            book=load_workbook(archive/"01_사용자자료"/"01_TMS"/"수질_SOOSIRO"/"SOOSIRO_수질TMS_정리.xlsx",read_only=True,data_only=True)
            self.assertIn("검토필요_연간",book.sheetnames)
            ws=book["검토필요_연간"]
            self.assertEqual(ws.max_row,2)
            headers=[c.value for c in next(ws.iter_rows(min_row=1,max_row=1))]
            self.assertIn("자료연도",headers)
            self.assertIn("원문 주소",headers)
            self.assertEqual(ws.cell(2,headers.index("SOOSIRO 사업장코드")+1).value,"31F0081")
            fidelity=json.loads((archive/"00_자료목록"/"Human_Delivery_Fidelity.json").read_text(encoding="utf-8"))
            self.assertTrue(fidelity["pass"])
            self.assertEqual(fidelity["sources"]["SOOSIRO_WATER"]["confirmed_rows"],0)
            self.assertEqual(fidelity["sources"]["SOOSIRO_WATER"]["review_rows"],2)
            self.assertFalse(fidelity["sources"]["SOOSIRO_WATER"]["silent_drop"])

    def test_prtr_and_chem_headers_preserve_source_semantics(self):
        with tempfile.TemporaryDirectory() as td:
            package=Path(td)/"assembled"; out=package/"output"; archive=package/"Human"
            self._csv(package/"Source_Identity.csv",[])
            for src in ["CLEANSYS_AIR","SOOSIRO_WATER"]:
                self._json(out/src/"status.json",{"status":"NO_MATCH"})
            self._jsonl(out/"CLEANSYS_AIR"/"annual_rows.jsonl",[]); self._json(out/"CLEANSYS_AIR"/"candidates.json",[])
            self._jsonl(out/"SOOSIRO_WATER"/"annual_rows.jsonl",[]); self._jsonl(out/"SOOSIRO_WATER"/"daily_rows.jsonl",[]); self._json(out/"SOOSIRO_WATER"/"fact_candidates.json",[])
            self._json(out/"PRTR"/"status.json",{"status":"DATA_FOUND"})
            self._csv(out/"PRTR"/"discovery.csv",[{"search_year":"2020","entrps_id":"300","company_name_raw":"효성티앤씨(주) 구미공장","address_raw":"경상북도 구미시 3공단2로 108","release_total_raw":"138","self_landfill_raw":"0","transfer_total_raw":"3181","source_url":"official"}])
            self._jsonl(out/"PRTR"/"detail_table_rows.jsonl",[{"search_year":2020,"entrps_id":"300","table_index":0,"row_index":0,"cells":["주 소","경상북도 구미시 3공단2로 108"]}])
            self._json(out/"CHEM_STATS"/"status.json",{"status":"DATA_FOUND"})
            self._csv(out/"CHEM_STATS"/"discovery.csv",[{"search_year":"2022","reportYear":"2022","bplcId":"AAW784N","bplcNm":"","locplcAdres":"","induty":"합성섬유 제조업","identity_anchor_year":"2020","identity_anchor_bplcNm":"효성티앤씨(주)구미공장","identity_anchor_locplcAdres":"경상북도 구미시 3공단2로 108"}])
            self._jsonl(out/"CHEM_STATS"/"detail_table_rows.jsonl",[{"search_year":2022,"bplcId":"AAW784N","table_index":0,"row_index":0,"cells":["업체명","효성티앤씨(주)구미공장","대표자","홍길동"]}])
            scope={"CLEANSYS_AIR":set(),"SOOSIRO_WATER":set(),"PRTR":{"300"},"CHEM_STATS":{"AAW784N"}}
            build_human_excels(package,archive,scope)

            prtr=load_workbook(archive/"01_사용자자료"/"02_화학물질"/"PRTR_배출이동량"/"PRTR_화학물질배출이동량_정리.xlsx",read_only=True,data_only=True)
            ph=[c.value for c in next(prtr["확정_사업장연도"].iter_rows(min_row=1,max_row=1))]
            self.assertEqual(ph[:4],["자료연도","원문 사업장명","원문 주소","PRTR 사업장ID"])
            dh=[c.value for c in next(prtr["확정_원문표"].iter_rows(min_row=1,max_row=1))]
            self.assertIn("원문 셀1",dh); self.assertNotIn("값1",dh)

            chem=load_workbook(archive/"01_사용자자료"/"02_화학물질"/"화학물질통계"/"화학물질통계_정리.xlsx",read_only=True,data_only=True)
            ch=[c.value for c in next(chem["확정_사업장연도"].iter_rows(min_row=1,max_row=1))]
            self.assertIn("해당연도 원문 주소",ch)
            self.assertIn("동일사업장 확인기준 주소",ch)
            ws=chem["확정_사업장연도"]
            self.assertIsNone(ws.cell(2,ch.index("해당연도 원문 주소")+1).value)
            self.assertEqual(ws.cell(2,ch.index("동일사업장 확인기준 주소")+1).value,"경상북도 구미시 3공단2로 108")


    def test_zero_row_index_and_stable_id_year_order(self):
        with tempfile.TemporaryDirectory() as td:
            package=Path(td)/"assembled"; out=package/"output"; archive=package/"Human"
            self._csv(package/"Source_Identity.csv",[])
            for src in ["CLEANSYS_AIR","SOOSIRO_WATER"]:
                self._json(out/src/"status.json",{"status":"NO_MATCH"})
            self._jsonl(out/"CLEANSYS_AIR"/"annual_rows.jsonl",[]); self._json(out/"CLEANSYS_AIR"/"candidates.json",[])
            self._jsonl(out/"SOOSIRO_WATER"/"annual_rows.jsonl",[]); self._jsonl(out/"SOOSIRO_WATER"/"daily_rows.jsonl",[]); self._json(out/"SOOSIRO_WATER"/"fact_candidates.json",[])

            self._json(out/"PRTR"/"status.json",{"status":"DATA_FOUND"})
            self._csv(out/"PRTR"/"discovery.csv",[
                {"search_year":"2021","entrps_id":"300","company_name_raw":"효성티앤씨(주) 구미공장","address_raw":"A"},
                {"search_year":"2020","entrps_id":"300","company_name_raw":"효성티앤씨(주)구미공장","address_raw":"A"},
            ])
            self._jsonl(out/"PRTR"/"detail_table_rows.jsonl",[
                {"search_year":2020,"entrps_id":"300","table_index":1,"row_index":1,"cells":["row1"]},
                {"search_year":2020,"entrps_id":"300","table_index":1,"row_index":0,"cells":["row0"]},
            ])

            self._json(out/"CHEM_STATS"/"status.json",{"status":"DATA_FOUND"})
            self._csv(out/"CHEM_STATS"/"discovery.csv",[
                {"search_year":"2022","bplcId":"AAW784N","bplcNm":"효성티앤씨(주) 구미공장","locplcAdres":"A"},
                {"search_year":"2020","bplcId":"AAW784N","bplcNm":"효성티앤씨(주)구미공장","locplcAdres":"A"},
                {"search_year":"2024","bplcId":"AAW784N","bplcNm":"효성티앤씨㈜ 구미공장","locplcAdres":"A"},
            ])
            self._jsonl(out/"CHEM_STATS"/"detail_table_rows.jsonl",[])
            scope={"CLEANSYS_AIR":set(),"SOOSIRO_WATER":set(),"PRTR":{"300"},"CHEM_STATS":{"AAW784N"}}
            build_human_excels(package,archive,scope)

            prtr=load_workbook(archive/"01_사용자자료"/"02_화학물질"/"PRTR_배출이동량"/"PRTR_화학물질배출이동량_정리.xlsx",read_only=True,data_only=True)
            ws=prtr["확정_원문표"]
            self.assertEqual(ws.cell(2,4).value,0)
            self.assertEqual(ws.cell(3,4).value,1)

            chem=load_workbook(archive/"01_사용자자료"/"02_화학물질"/"화학물질통계"/"화학물질통계_정리.xlsx",read_only=True,data_only=True)
            ws=chem["확정_사업장연도"]
            years=[ws.cell(r,1).value for r in range(2,5)]
            self.assertEqual(years,["2020","2022","2024"])


if __name__ == "__main__":
    unittest.main()
